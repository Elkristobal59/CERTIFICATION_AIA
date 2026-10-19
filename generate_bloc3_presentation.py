import os
import sys

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")

from pptx import Presentation

from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

OUTPUT_DIR = os.path.dirname(os.path.abspath(__file__))
PPTX_PATH = os.path.join(OUTPUT_DIR, "ARCHITECTURE_PIPELINE_FRAUD_BLOC3.pptx")

# Palette graphique premium
NAVY = RGBColor(15, 23, 42)        # Fond sombre / Header
SLATE = RGBColor(30, 41, 59)       # Carte
BLUE = RGBColor(14, 165, 233)       # Accent Cyan/Bleu
GREEN = RGBColor(16, 185, 129)     # Succès
RED = RGBColor(239, 68, 68)        # Alerte
WHITE = RGBColor(255, 255, 255)    # Texte blanc
GRAY = RGBColor(148, 163, 184)     # Texte secondaire
LIGHT_BG = RGBColor(248, 250, 252) # Fond clair diapo

def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    def add_header(slide, title_text, category_text="CERTIFICATION ARCHITECTE IA — RNCP 41993 — BLOC 3"):
        # Bandeau de fond
        header_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(1.15))
        header_shape.fill.solid()
        header_shape.fill.fore_color.rgb = NAVY
        header_shape.line.fill.background()

        # Catégorie
        txbox = slide.shapes.add_textbox(Inches(0.8), Inches(0.12), Inches(11.7), Inches(0.35))
        tf = txbox.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = category_text.upper()
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = BLUE

        # Titre
        txbox_t = slide.shapes.add_textbox(Inches(0.8), Inches(0.42), Inches(11.7), Inches(0.65))
        tf_t = txbox_t.text_frame
        tf_t.word_wrap = True
        p_t = tf_t.paragraphs[0]
        p_t.text = title_text
        p_t.font.size = Pt(22)
        p_t.font.bold = True
        p_t.font.color.rgb = WHITE

    # ==========================================
    # SLIDE 1 : COUVERTURE
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = NAVY
    bg1.line.fill.background()

    tb1 = s1.shapes.add_textbox(Inches(1.2), Inches(1.6), Inches(11.0), Inches(4.5))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "CERTIFICATION ARCHITECTE EN INTELLIGENCE ARTIFICIELLE"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = BLUE
    p.space_after = Pt(15)

    p = tf1.add_paragraph()
    p.text = "PIPELINE DE STREAMING & DÉTECTION DE FRAUDE EN TEMPS RÉEL"
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.space_after = Pt(20)

    p = tf1.add_paragraph()
    p.text = "Architecture Hybride : Apache Kafka KRaft • GCP Cloud Run / Docker • Airflow • Parquet Lake"
    p.font.size = Pt(18)
    p.font.color.rgb = GRAY
    p.space_after = Pt(35)

    p = tf1.add_paragraph()
    p.text = "Bloc de Compétences 3 (RNCP 41993) : Concevoir et mettre en œuvre des pipelines de données pour l'IA\nCandidat : Chris GILLERON | Lead Data & AI Architect"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = GREEN

    # ==========================================
    # SLIDE 2 : CONTEXTE MÉTIER & PROBLÉMATIQUE
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "1. Enjeux Stratégiques, Problématique & Exigences Métier")

    cards_data = [
        ("🚨 Le Défi Financier", "1,00+ Milliard €", "Montant des transactions frauduleuses par carte en UE (BCE). Les banques et fintechs subissent une vélocité accrue des attaques nécessitant une réponse en millisecondes."),
        ("⚡ Double Exigence Métier", "Temps Réel + Batch", "1. Notification immédiate dès qu'un paiement suspect survient pour blocage préventif.\n2. Rapport d'audit matinal automatisé (06h00) récapitulant les flux des 24h."),
        ("🛡️ Contraintes Réglementaires", "PCI-DSS & RGPD", "Interdiction absolue de stocker des numéros de carte (PAN) en clair. Pseudonymisation des données clients et isolation immédiate des messages non conformes (DLQ).")
    ]
    for idx, (head, metric, desc) in enumerate(cards_data):
        x = Inches(0.8 + idx * 4.0)
        card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.8), Inches(3.7), Inches(5.0))
        card.fill.solid()
        card.fill.fore_color.rgb = SLATE
        card.line.color.rgb = BLUE
        card.line.width = Pt(1.5)

        tx = s2.shapes.add_textbox(x + Inches(0.2), Inches(2.0), Inches(3.3), Inches(4.5))
        tf = tx.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = head
        p.font.size = Pt(18)
        p.font.bold = True
        p.font.color.rgb = WHITE
        p.space_after = Pt(14)

        p = tf.add_paragraph()
        p.text = metric
        p.font.size = Pt(22)
        p.font.bold = True
        p.font.color.rgb = BLUE
        p.space_after = Pt(16)

        p = tf.add_paragraph()
        p.text = desc
        p.font.size = Pt(13)
        p.font.color.rgb = GRAY

    # ==========================================
    # SLIDE 3 : ARCHITECTURE TECHNIQUE DU PIPELINE
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, "2. Architecture Technique de Bout en Bout (Hybride Cloud / On-Premise)")

    pipeline_steps = [
        ("1. Ingestion Live", "API HuggingFace &\nFlux Replay Sécurisé\nVélocité continue", BLUE),
        ("2. Stream Broker", "Apache Kafka KRaft\nTopic 'raw-transactions'\n3 partitions équilibrées", GREEN),
        ("3. Inférence ML", "FastAPI Conteneurisé\nGCP Cloud Run / Docker\nScale-to-Zero, < 15ms", BLUE),
        ("4. Alerting & Sécurité", "Console, Logs & DB\nMasquage PCI-DSS\nDécision auto 3DS/Block", RED),
        ("5. Dual Storage", "Data Lake Bronze Parquet\n+ SQLite / PostgreSQL\nAudit Trail immuable", GREEN),
        ("6. Batch Audit", "Apache Airflow DAG\nDéclenchement 06h00 UTC\nRapport Matinal Markdown", BLUE)
    ]
    for idx, (title, details, col) in enumerate(pipeline_steps):
        row = 0 if idx < 3 else 1
        col_idx = idx % 3
        x = Inches(0.8 + col_idx * 4.0)
        y = Inches(1.8 + row * 2.6)

        box = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(3.7), Inches(2.3))
        box.fill.solid()
        box.fill.fore_color.rgb = SLATE
        box.line.color.rgb = col
        box.line.width = Pt(2)

        tx = s3.shapes.add_textbox(x + Inches(0.2), y + Inches(0.15), Inches(3.3), Inches(2.0))
        tf = tx.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = col
        p.space_after = Pt(10)

        p = tf.add_paragraph()
        p.text = details
        p.font.size = Pt(12)
        p.font.color.rgb = WHITE

    # ==========================================
    # SLIDE 4 : CONTRÔLE QUALITÉ & DEAD LETTER QUEUE (C5)
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, "3. Contrôle Qualité des Flux, Résilience & Dead Letter Queue (C5)")

    col1 = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0))
    col1.fill.solid()
    col1.fill.fore_color.rgb = SLATE
    col1.line.color.rgb = BLUE
    col1.line.width = Pt(1.5)

    tx1 = s4.shapes.add_textbox(Inches(1.1), Inches(2.1), Inches(5.0), Inches(4.4))
    tf1 = tx1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "🎯 Validation de Schéma & Data Quality Gate"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.space_after = Pt(14)

    p = tf1.add_paragraph()
    p.text = "• Vérification Pydantic des attributs obligatoires (montant, coordonnées géographiques, date/heure).\n\n• Règles Métier Strictes : rejet immédiat des montants négatifs (exploits et failles de débordement).\n\n• Métriques de qualité : calcul de complétude, vérification de types et détection des valeurs aberrantes avant scoring."
    p.font.size = Pt(13)
    p.font.color.rgb = GRAY

    col2 = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.9), Inches(1.8), Inches(5.6), Inches(5.0))
    col2.fill.solid()
    col2.fill.fore_color.rgb = SLATE
    col2.line.color.rgb = RED
    col2.line.width = Pt(1.5)

    tx2 = s4.shapes.add_textbox(Inches(7.2), Inches(2.1), Inches(5.0), Inches(4.4))
    tf2 = tx2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "🚨 Mécanisme de Dead Letter Queue (DLQ)"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.space_after = Pt(14)

    p = tf2.add_paragraph()
    p.text = "• Topic Dédié 'fraud-dlq' : les messages corrompus ne bloquent jamais la file principale des transactions saines.\n\n• Isolation en Base : table SQL `dead_letter_queue` enregistrant le payload brut et la cause d'erreur exacte.\n\n• Résilience Opérationnelle : permet le rejeu asynchrone des transactions après correction de schéma ou de bug."
    p.font.size = Pt(13)
    p.font.color.rgb = GRAY

    # ==========================================
    # SLIDE 5 : SÉCURITÉ, PCI-DSS & RGPD (C4)
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, "4. Sécurité, Confidentialité des Données & Conformité PCI-DSS (C4)")

    sec_cards = [
        ("💳 Masquage PCI-DSS", "Cartes Masquées 6+4", "Le numéro de carte (PAN) est tronqué : seuls les 6 premiers chiffres (BIN) et les 4 derniers sont conservés (`453288******9812`). Aucun numéro complet en base."),
        ("👤 Pseudonymisation RGPD", "Privacy-by-Design", "Les noms des porteurs sont hachés via SHA-256 et masqués (`J*** D*** ID:A3F8`). Respect intégral des principes de minimisation des données (Article 5 RGPD)."),
        ("⚡ SLA & Alertes Immédiates", "< 15 ms Latence", "Modèle HistGradientBoosting optimisé. Alerte visuelle console et consignation en log immuable dès probabilité >= 50% avec classification de risque (CRITICAL, HIGH).")
    ]
    for idx, (title, sub, text) in enumerate(sec_cards):
        x = Inches(0.8 + idx * 4.0)
        card = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.8), Inches(3.7), Inches(5.0))
        card.fill.solid()
        card.fill.fore_color.rgb = SLATE
        card.line.color.rgb = GREEN
        card.line.width = Pt(1.5)

        tx = s5.shapes.add_textbox(x + Inches(0.2), Inches(2.0), Inches(3.3), Inches(4.5))
        tf = tx.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(18)
        p.font.bold = True
        p.font.color.rgb = WHITE
        p.space_after = Pt(12)

        p = tf.add_paragraph()
        p.text = sub
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = GREEN
        p.space_after = Pt(14)

        p = tf.add_paragraph()
        p.text = text
        p.font.size = Pt(13)
        p.font.color.rgb = GRAY

    # ==========================================
    # SLIDE 6 : INDUSTRIALISATION, FINOPS & SYNTHÈSE
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    add_header(s6, "5. Synthèse des Livrables & Démonstration du Pipeline en Production")

    c_box = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
    c_box.fill.solid()
    c_box.fill.fore_color.rgb = SLATE
    c_box.line.color.rgb = BLUE
    c_box.line.width = Pt(1.5)

    tx_s = s6.shapes.add_textbox(Inches(1.2), Inches(2.1), Inches(11.0), Inches(4.4))
    tf_s = tx_s.text_frame
    tf_s.word_wrap = True

    p = tf_s.paragraphs[0]
    p.text = "🎯 Validation Intégrale des Attentes du Jury (RNCP 41993 - Bloc 3)"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.space_after = Pt(16)

    bullets = [
        ("📁 Dépôt GitHub Officiel", "Code versionné sur https://github.com/Elkristobal59/CERTIFICATION_AIA avec CI/CD Cloud Build et .gitignore rigoureux."),
        ("☁️ Déploiement Serverless Cloud Run", "Microservice FastAPI déployable avec Scale-to-Zero (min-instances: 0, facturation nulle hors trafic, dimensionné à 1 vCPU / 1 GiB)."),
        ("🎥 Vidéo de Démonstration Live", "Enregistrement pas-à-pas démontrant le flux Kafka, l'alerte en direct, le routage DLQ et l'exécution du rapport matinal Airflow."),
        ("📊 Modèle & Tracking MLflow", "Modèle supervisé avec métriques logguées dans MLflow (ROC-AUC: 0.98+, Rappel: 92%+, seuil calibré) et Model Registry."),
        ("📈 Double Temporalité Métier", "Exigence 1 (Alerte live < 15ms) et Exigence 2 (Batch matinal d'audit 06h00) validées en environnement réel conteneurisé.")
    ]
    for b_title, b_desc in bullets:
        p = tf_s.add_paragraph()
        p.text = f"• {b_title} : {b_desc}"
        p.font.size = Pt(13)
        p.font.color.rgb = WHITE
        p.space_after = Pt(8)

    prs.save(PPTX_PATH)
    print(f"✅ Présentation PowerPoint générée avec succès : {PPTX_PATH}")

if __name__ == "__main__":
    create_deck()
